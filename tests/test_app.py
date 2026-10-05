from io import BytesIO
from pathlib import Path
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from model import DEFAULTS, METHODS
from streamlit_lackverbrauch_app import export_data, export_images, CHARTS


class AppTests(unittest.TestCase):
    def test_routes_snapshot_compare_invalid_and_reset(self):
        app = AppTest.from_file(str(Path(__file__).parents[1] / 'streamlit_lackverbrauch_app.py')).run(timeout=30)
        self.assertFalse(app.exception)
        app.text_input(key='scenario_name').set_value('Bekannt')
        next(b for b in app.button if b.label == 'Szenario übernehmen').click().run(timeout=30)
        self.assertEqual(len(app.session_state['scenarios']), 1)
        app.radio(key='method').set_value(METHODS[1]).run(timeout=30)
        app.number_input(key='mng').set_value(70.0)
        app.text_input(key='scenario_name').set_value('TDS')
        next(b for b in app.button if b.label == 'Szenario übernehmen').click().run(timeout=30)
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state['scenarios'][0]['params']['mng'], 50)
        self.assertEqual(len(app.session_state['scenarios']), 2)
        app.radio(key='method').set_value(METHODS[2]).run(timeout=30)
        self.assertFalse(app.exception)
        app.number_input(key='pause').set_value(100.0).run(timeout=30)
        self.assertTrue(app.error)
        self.assertFalse(app.exception)
        self.assertTrue(next(b for b in app.button if b.label == 'Szenario übernehmen').disabled)
        next(b for b in app.button if b.label == 'Reset').click().run(timeout=30)
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state['scenarios'], [])
        self.assertEqual(app.number_input(key='pause').value, 10)

    def test_exports_and_chart_endpoints(self):
        entries = [dict(name='Referenz', params=DEFAULTS.copy(), color='#0072B2')]
        csv, excel = export_data(entries)
        self.assertIn('Lackverlust [L]', csv.decode('utf-8-sig'))
        workbook = load_workbook(BytesIO(excel))
        self.assertEqual(workbook.sheetnames, ['Zeiträume', 'Parameter', 'Verläufe', 'Einheiten'])
        self.assertEqual(workbook['Zeiträume'].max_row, 5)
        for chart in CHARTS:
            png, svg = export_images(entries, chart, 'Jahr')
            self.assertTrue(png.startswith(b'\x89PNG'))
            self.assertIn(b'<svg', svg)


if __name__ == '__main__':
    unittest.main()
